#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown (Persian, RTL) -> DOCX with proper bidi, embedded-table handling.

Usage: python3 build_docx.py <source.md> <out.docx>
Requires: python-docx
"""
import re
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1]
OUT = sys.argv[2]
FONT = "Vazirmatn"
ACCENT = RGBColor(0x0B, 0x3D, 0x38)
INLINE = re.compile(r"(\*\*.+?\*\*|(?<![\*\w])\*(?!\s).+?(?<!\s)\*(?![\*\w]))")


def set_rtl(p):
    p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))


def style_run(run, size=None, bold=None, color=None):
    run.font.name = FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts")) or OxmlElement("w:rFonts")
    if rFonts.getparent() is None:
        rPr.append(rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(a), FONT)
    rtl = OxmlElement("w:rtl"); rtl.set(qn("w:val"), "1"); rPr.append(rtl)
    if size:
        run.font.size = Pt(size)
        sz = OxmlElement("w:szCs"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_runs(p, text, size=11.5, color=None):
    for piece in INLINE.split(text):
        if not piece:
            continue
        bold = italic = False
        if piece.startswith("**") and piece.endswith("**"):
            piece, bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            piece, italic = piece[1:-1], True
        r = p.add_run(piece)
        style_run(r, size=size, bold=bold, color=color)
        if italic:
            r.italic = True


def shade(cell, fill="EAF1EF"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec._sectPr.append(OxmlElement("w:bidi"))

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11.5)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts")) or OxmlElement("w:rFonts")
    if rFonts.getparent() is None:
        rPr.append(rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(a), FONT)
    normal.element.get_or_add_pPr().append(OxmlElement("w:bidi"))

    for name, size in (("Heading 1", 18), ("Heading 2", 14.5), ("Heading 3", 12.5), ("Heading 4", 11.5)):
        st = doc.styles[name]
        st.font.name = FONT; st.font.size = Pt(size)
        st.font.color.rgb = ACCENT; st.font.bold = True
        st.paragraph_format.space_before = Pt(0 if name == "Heading 1" else 14)
        st.paragraph_format.space_after = Pt(6)

    lines = open(SRC, encoding="utf-8").read().split("\n")
    i, first_h1 = 0, False

    while i < len(lines):
        raw = lines[i].rstrip()

        if raw.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-\|]+\|$", lines[i + 1].strip()):
            head = [c.strip() for c in raw.strip().strip("|").split("|")]
            rows, j = [], i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            table = doc.add_table(rows=1, cols=len(head))
            table.style = "Table Grid"
            for k, txt in enumerate(head):
                p = table.rows[0].cells[k].paragraphs[0]
                set_rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                add_runs(p, txt, size=9, color=ACCENT)
                shade(table.rows[0].cells[k])
            for r in rows:
                cells = table.add_row().cells
                for k, txt in enumerate(r[:len(head)]):
                    p = cells[k].paragraphs[0]
                    set_rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    add_runs(p, txt, size=9)
            doc.add_paragraph()
            i = j
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", raw)
        if m:
            level, txt = len(m.group(1)), m.group(2).strip()
            if level == 1 and not first_h1:
                p = doc.add_paragraph(); set_rtl(p)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, txt, size=20, color=ACCENT)
                for r in p.runs:
                    r.bold = True
                first_h1 = True
            else:
                p = doc.add_paragraph(style="Heading %d" % min(level, 4))
                set_rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                add_runs(p, txt, size={2: 14.5, 3: 12.5, 4: 11.5}.get(level, 12.5), color=ACCENT)
                for r in p.runs:
                    r.bold = True
            i += 1
            continue

        if re.match(r"^---+\s*$", raw):
            i += 1
            continue

        if raw.startswith(">"):
            txt = raw[1:].strip()
            if txt:
                p = doc.add_paragraph(); set_rtl(p)
                p.paragraph_format.right_indent = Cm(0.8)
                p.paragraph_format.left_indent = Cm(0.8)
                add_runs(p, txt, size=11, color=RGBColor(0x16, 0x33, 0x2F))
            i += 1
            continue

        mol = re.match(r"^([۰-۹0-9]+)[.)]\s+(.*)$", raw)
        mul = re.match(r"^[-*]\s+(.*)$", raw)
        if mol or mul:
            p = doc.add_paragraph(); set_rtl(p)
            p.paragraph_format.right_indent = Cm(0.7)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, (mol.group(1) + ". " + mol.group(2)) if mol else ("• " + mul.group(1)))
            i += 1
            continue

        if not raw.strip():
            i += 1
            continue

        buf, j = [raw.strip()], i + 1
        while j < len(lines):
            nxt = lines[j].rstrip()
            if (not nxt.strip() or nxt.startswith("#") or nxt.startswith(">") or nxt.startswith("|")
                    or re.match(r"^---+\s*$", nxt) or re.match(r"^[-*]\s+", nxt)
                    or re.match(r"^([۰-۹0-9]+)[.)]\s+", nxt)):
                break
            buf.append(nxt.strip())
            j += 1
        p = doc.add_paragraph(); set_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.45
        add_runs(p, " ".join(buf))
        i = j

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
